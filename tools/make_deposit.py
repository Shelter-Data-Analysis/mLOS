"""Stage a Zenodo deposit: what travels, its digests, and what reads it.

Two bundles, because the results and the slide deck have different half-lives
and only one of them is what a paper cites.  See PUBLISHING.md.

    python3 tools/make_deposit.py            # the results bundle
    python3 tools/make_deposit.py --deck     # the deck bundle

Each lands beside the repository by default, in `_mlos_results` and
`_mlos_deck`, so both sit outside git.  A destination may be given instead.

Run after a regeneration pass, once every run has been made under the version
being tagged.  A bundle whose runs disagree about the version cannot say what
produced it, and that is the whole point of depositing it, so the mixed case
is refused rather than reported.

Provenance is read out of each run's `results.json` rather than typed here, so
the bundle records what a run actually did.

Tracked, though nothing a run does reads it: what a deposit contains, and
what disqualifies a file from one, is a rule that has to outlive the laptop it
was first enforced on.  PUBLISHING.md is the prose; this is the enforcement.
"""

from __future__ import annotations

import argparse
import csv
import datetime
import hashlib
import json
import re
import shutil
import sys
import textwrap
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
RESULTS_OUT = ROOT.parent / "_mlos_results"
DECK_OUT = ROOT.parent / "_mlos_deck"

#  The version DOI of the prepared data every shipped run reads.  A version
#  DOI rather than a concept one: a digest pins bytes, and only the version
#  record is guaranteed to still hold them.
PREPARED_DOI = "10.5281/zenodo.22051368"

#  mLOS's own CONCEPT DOI, which resolves to the newest release.  It does not
#  exist until the first release is archived; set it then, and re-run.  The
#  README names the version alongside it, and that pair is what a reader
#  checks against a run log.
SOFTWARE_DOI = "10.5281/zenodo.22083814"

#  The results deposit's version DOI, needed only by the deck bundle, which
#  is derived from it.  Set it once the results deposit is published.
RESULTS_DOI = "10.5281/zenodo.22652165"

#  What each run is, and how far it has been checked.  A deposited result
#  without this is a number with no provenance for its trustworthiness.
DATASET = {
    "OC2": ("Orange County Animal Care, all species, four intake types",
            "the current definition of the dataset, and the default example"
            " the guides quote"),
    "OC1": ("Orange County Animal Care, dogs",
            "the previous definition of the same dataset, frozen and kept as"
            " a baseline to check against"),
}

#  Which deck belongs to which run.  The pairing is by hand because nothing in
#  a .pptx says which results.json it was built from; the script checks the
#  pairing it is given rather than inferring one.
DECKS = {
    "OC2": ("mlos_deck.pptx", "mlos_deck_tables.xlsx", "mlos_deck_figures"),
    "OC1": ("mlos_deck_OC1.pptx", "mlos_deck_OC1_tables.xlsx",
            "mlos_deck_OC1_figures"),
}

#  Variant decks deposited beside the main one, per run.  Each carries its own
#  figures, at <stem>_figures, and no workbook: a variant's workbook would be
#  the deck's under another name.  Branded builds stay out, because the
#  template is not in the repository and a deposit nobody can rebuild from the
#  repository is the one link in this chain that would not hold; so does
#  extended_variant_features, which says in its own first lines that it is
#  written to be read beside presentation_guide.md rather than shown to a room.
#  PUBLISHING.md step 7 carries the reasoning.
VARIANT_DECKS = {
    "OC2": ("educational.pptx",),
    "OC1": (),
}

#  The sidecar mlos_review.deck writes beside a deck, naming every slide's
#  position, title, layout and run.  It travels with the deck because that is
#  what mlos_review.variant reads to build an outline against one, so a deposit
#  without it ships a deck that can be shown and not extended.  Variants write
#  none of their own; there is one per main deck.
SLIDES_SUFFIX = "_slides.json"

#  What each file at the top of a run's folder is called in the manifest.
DECK_KINDS = {".pptx": "deck", ".xlsx": "table workbook",
              ".json": "slide manifest"}

#  Order runs appear in the manifest and the README.
ORDER = ["OC2", "OC1"]

#  Runs that are computed but not deposited.  OC1 is the previous definition
#  of the same dataset, kept and regenerated as a baseline the guides quote and
#  tests/show_guide_examples.py checks, but it is not what the analysis reports
#  on.  Its input and its settings file are published already, in the prepared
#  data deposit and in this repository, so nothing about it is unavailable;
#  what is not deposited is a second set of results nothing cites.  Skipped
#  before the version check, so declining to regenerate OC1 cannot block an
#  OC2 deposit.
SKIP = {"OC1"}

#  Left out of a run directory: a scratch workbook the test suite writes, and
#  the Finder's leavings.
SKIP_FILES = {"from_json.xlsx", ".DS_Store"}

KINDS = [
    ("results.json", "results bundle"),
    ("analysis_results.xlsx", "Excel workbook"),
    ("analysis_log.txt", "console log"),
    ("data_preparation_stats.csv", "preparation statistics"),
]


def sha256(path):
    digest = hashlib.sha256()
    with open(str(path), "rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def kind_of(name):
    for exact, label in KINDS:
        if name == exact:
            return label
    if name.endswith(".png"):
        return "plot"
    if name.endswith(".csv"):
        return "curve values"
    if name.endswith((".yaml", ".yml")):
        return "settings file"
    if name.endswith("_data.md"):
        return "input provenance"
    if name.endswith("_data_run.txt"):
        return "input preparation log"
    return "other"


def prep_field(log, label):
    found = re.search(r"^{0}\s+(.+?)\s*$".format(re.escape(label)), log, re.M)
    return found.group(1) if found else ""


def read_runs():
    """One record per run directory, straight out of what the run recorded.

    A run directory is `results/` itself or a directory one level under it
    holding a `results.json`.  The archive directories a run leaves behind are
    named `mLOS_<date>_<n>` and are skipped: they hold the previous run, which
    is the version this pass exists to stop depositing.
    """
    results = ROOT / "results"
    if not results.is_dir():
        sys.exit("no results/ directory -- run the analysis first")
    candidates = [results] + sorted(
        path for path in results.iterdir()
        if path.is_dir() and not path.name.startswith("mLOS_"))

    runs = []
    for path in candidates:
        bundle_file = path / "results.json"
        if not bundle_file.is_file():
            continue
        with open(str(bundle_file), encoding="utf-8") as handle:
            bundle = json.load(handle)
        run = bundle.get("run", {})
        label = re.sub(r"_data\.csv$", "", Path(run.get("data_file", "")).name)
        if label in SKIP:
            continue
        version = run.get("mlos_version")
        if not version:
            sys.exit(
                "{0} carries no mlos_version -- it predates version stamping, "
                "so rerun that analysis before depositing".format(
                    bundle_file.relative_to(ROOT)))
        data_file = ROOT / run["data_file"]
        runs.append({
            "label": label,
            "dir": path,
            "version": version,
            "generated_at": run["generated_at"],
            "schema_version": bundle.get("schema_version"),
            "data_file": run["data_file"],
            "data_path": data_file,
            "recorded_data_sha256": run.get("data_sha256"),
            "settings_file": run["settings_file"],
            "settings_path": ROOT / run["settings_file"],
            "recorded_settings_sha256": run.get("settings_sha256"),
        })

    if not runs:
        sys.exit("no results.json under results/ -- run the analysis first")

    versions = sorted({run["version"] for run in runs})
    if len(versions) > 1:
        sys.exit("results/ mixes versions {0} -- regenerate every run under "
                 "one before depositing".format(", ".join(versions)))

    unknown = [run["label"] for run in runs if run["label"] not in DATASET]
    if unknown:
        sys.exit("no DATASET entry for {0} -- add it, and add it to ORDER, in "
                 "this script".format(", ".join(sorted(unknown))))

    duplicates = sorted({run["label"] for run in runs
                         if [r["label"] for r in runs].count(run["label"]) > 1})
    if duplicates:
        sys.exit("two run directories report the same dataset {0} -- one of "
                 "them is stale".format(", ".join(duplicates)))

    runs.sort(key=lambda run: ORDER.index(run["label"])
              if run["label"] in ORDER else len(ORDER))
    return runs


def check_input(run):
    """Confirm the input is the deposited byte stream, not a local variant.

    The preparation log beside the input records the digest ShelterDataPrep
    wrote.  Recomputing it here is what lets the deposit claim the prepared
    data DOI as its source rather than merely naming it.

    Then the run's own record, which closes the other gap: the prep log says
    what was prepared and this pass says what is about to be staged, and
    neither says what the analysis read.  A file edited between the run and
    the deposit satisfies both of those and none of the numbers.
    """
    prep_log = run["data_path"].with_name(
        run["data_path"].name.replace(".csv", "_run.txt"))
    if not run["data_path"].is_file():
        sys.exit("{0} is missing -- the run names it as its input".format(
            run["data_file"]))
    if not prep_log.is_file():
        sys.exit(
            "{0} has no preparation log beside it, so the deposit cannot show "
            "its input is the deposited one; copy it over from "
            "ShelterDataPrep's results/".format(run["data_file"]))
    log = prep_log.read_text(encoding="utf-8")
    recorded = prep_field(log, "output sha256")
    actual = sha256(run["data_path"])
    if recorded != actual:
        sys.exit(
            "{0} does not match the digest in {1}:\n  log   {2}\n  file  {3}\n"
            "The input is not the file that log describes.".format(
                run["data_file"], prep_log.name, recorded or "(absent)",
                actual))
    run["data_sha256"] = actual
    run["prep_log"] = prep_log
    run["prep_version"] = prep_field(log, "shelterprep")
    check_run_digests(run)
    return run


def check_run_digests(run):
    """The two files the run recorded, against the two about to be staged.

    A run that predates digest stamping is refused the way one predating
    version stamping is: the point of the deposit is to say what produced the
    numbers, and rerunning is cheap.
    """
    for kind in ("data", "settings"):
        recorded = run["recorded_{0}_sha256".format(kind)]
        path = run["{0}_path".format(kind)]
        if not recorded:
            sys.exit(
                "{0} carries no {1}_sha256 -- it predates digest stamping, so "
                "rerun that analysis before depositing".format(
                    run["dir"].relative_to(ROOT), kind))
        if recorded.startswith("("):
            sys.exit(
                "{0} recorded '{1}' for its {2} file instead of a digest; "
                "install the digest R package and rerun that analysis".format(
                    run["dir"].relative_to(ROOT), recorded, kind))
        actual = sha256(path)
        if recorded != actual:
            sys.exit(
                "{0} is not the file the run read:\n  run   {1}\n  file  {2}\n"
                "The deposit would ship an input the numbers did not come "
                "from.".format(run["{0}_file".format(kind)], recorded, actual))


def fresh(out):
    if out.exists():
        shutil.rmtree(str(out))
    out.mkdir(parents=True)


def write_manifest(out, header, rows):
    with open(str(out / "MANIFEST.csv"), "w", newline="",
              encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def stage_results(runs, out):
    """Copy what travels, and return a manifest row for each file."""
    fresh(out)
    rows = []
    for run in runs:
        folder = out / run["label"]
        folder.mkdir()
        extras = [run["settings_path"], run["prep_log"],
                  run["data_path"].with_suffix(".md")]
        missing = [path.name for path in extras if not path.is_file()]
        if missing:
            sys.exit("{0}: {1} not found -- a deposited run travels with the "
                     "settings that produced it and the provenance of its "
                     "input".format(run["label"], ", ".join(missing)))
        travels = sorted(path for path in run["dir"].iterdir()
                         if path.is_file() and path.name not in SKIP_FILES)
        for source in travels + extras:
            shutil.copy2(str(source), str(folder / source.name))
            rows.append([
                "{0}/{1}".format(run["label"], source.name),
                kind_of(source.name), run["label"], run["version"],
                run["generated_at"], sha256(source), source.stat().st_size,
                run["data_file"], run["data_sha256"], run["settings_file"]])
    rows.sort(key=lambda row: (ORDER.index(row[2]) if row[2] in ORDER
                               else len(ORDER), row[1], row[0]))
    write_manifest(out, ["file", "kind", "dataset", "mlos_version",
                         "generated_at", "sha256", "bytes", "input_file",
                         "input_sha256", "settings_file"], rows)
    return rows


DECK_PROVENANCE = re.compile(
    r"Statistics generated on (\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})")


def deck_run_timestamp(path):
    """The run timestamp a deck prints on its opening slide, or None.

    An mtime says a file is newer, which a rebuild against the wrong run
    satisfies as easily as one against the right run.  The printed timestamp
    is the deck's own claim about which analysis it shows, so it is the one
    worth checking against results.json.
    """
    from pptx import Presentation  # only the deck bundle needs it

    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            found = DECK_PROVENANCE.search(shape.text_frame.text)
            if found:
                return found.group(1)
    return None


def check_deck_matches_run(deck, run, bundle_file):
    """Refuse a deck that is older than its run, or built from a different one."""
    if deck.stat().st_mtime < bundle_file.stat().st_mtime:
        sys.exit(
            "{0} is older than the run it claims to show ({1}) -- rebuild "
            "the deck".format(deck.relative_to(ROOT),
                              bundle_file.relative_to(ROOT)))
    printed = deck_run_timestamp(deck)
    expected = run["generated_at"][:19]
    if printed is None:
        sys.exit(
            "{0} prints no run timestamp, so nothing in it says which "
            "analysis it shows".format(deck.relative_to(ROOT)))
    if printed != expected:
        sys.exit(
            "{0} shows a different run:\n  deck  {1}\n  run   {2}\n"
            "Rebuild it from the run being deposited.".format(
                deck.relative_to(ROOT), printed, expected))


def deck_file_kind(name, suffix):
    """What a staged deck file is called in the manifest.

    Nesting decides it before the extension does.  A figures directory carries
    its own manifest.json beside the images, naming what was drawn; on
    extension alone that reads as the deck's slide manifest, which is a
    different file answering a different question.
    """
    if "/" in name:
        return ("figure manifest" if name.endswith("/manifest.json")
                else "figure")
    return DECK_KINDS.get(suffix, "figure")


def stage_deck(runs, out):
    """Copy each deck with the workbook and figures it was built from."""
    fresh(out)
    rows = []
    settings = ROOT / "data" / "OC_deck_settings.yaml"
    if not settings.is_file():
        sys.exit("data/OC_deck_settings.yaml not found -- the deck settings "
                 "are how a reader knows what the deck selected")
    for run in runs:
        if run["label"] not in DECKS:
            sys.exit("no DECKS entry for {0} -- add it, or drop that run from "
                     "the deck bundle".format(run["label"]))
        deck_name, tables_name, figures_name = DECKS[run["label"]]
        deck = ROOT / "reports" / deck_name
        tables = ROOT / "reports" / tables_name
        figures = ROOT / "reports" / figures_name
        # Derived from the deck's own name, the rule mlos_review.deck writes it
        # under, so a renamed deck cannot leave its manifest behind.  A variant
        # is built against this file and refuses without it, so a deck deposited
        # without one is a deck nobody downstream can write a variant for.
        slides = deck.with_name(deck.stem + SLIDES_SUFFIX)
        for path in (deck, tables, figures, slides):
            if not path.exists():
                sys.exit("{0} not found -- rebuild the deck for {1} before "
                         "depositing it".format(
                             path.relative_to(ROOT), run["label"]))
        bundle_file = run["dir"] / "results.json"
        check_deck_matches_run(deck, run, bundle_file)

        folder = out / run["label"]
        folder.mkdir()
        for source in (deck, tables, slides):
            shutil.copy2(str(source), str(folder / source.name))
        shutil.copytree(str(figures), str(folder / "figures"))
        staged = [folder / deck.name, folder / tables.name,
                  folder / slides.name]
        staged += sorted(path for path in (folder / "figures").iterdir()
                         if path.is_file() and path.name not in SKIP_FILES)

        # The variants, each with its own figures.  Named rather than
        # discovered: reports/ also holds the branded builds and whatever was
        # last tried out, and a deposit that sweeps a directory deposits those
        # too, one release after somebody experimented.
        for variant_name in VARIANT_DECKS.get(run["label"], ()):
            variant = ROOT / "reports" / variant_name
            variant_figures = variant.with_name(variant.stem + "_figures")
            for path in (variant, variant_figures):
                if not path.exists():
                    sys.exit(
                        "{0} not found -- rebuild the {1} variant for {2} "
                        "before depositing it".format(
                            path.relative_to(ROOT), variant.stem,
                            run["label"]))
            check_deck_matches_run(variant, run, bundle_file)
            shutil.copy2(str(variant), str(folder / variant.name))
            shutil.copytree(str(variant_figures),
                            str(folder / variant_figures.name))
            staged.append(folder / variant.name)
            staged += sorted(
                path for path in (folder / variant_figures.name).iterdir()
                if path.is_file() and path.name not in SKIP_FILES)
        for source in staged:
            name = str(source.relative_to(folder)).replace("\\", "/")
            rows.append([
                "{0}/{1}".format(run["label"], name),
                deck_file_kind(name, source.suffix),
                run["label"], run["version"], run["generated_at"],
                sha256(source), source.stat().st_size])
    shutil.copy2(str(settings), str(out / settings.name))
    rows.append([settings.name, "deck settings", "", runs[0]["version"], "",
                 sha256(settings), settings.stat().st_size])
    rows.sort(key=lambda row: (ORDER.index(row[2]) if row[2] in ORDER
                               else len(ORDER), row[1], row[0]))
    write_manifest(out, ["file", "kind", "dataset", "mlos_version",
                         "generated_at", "sha256", "bytes"], rows)
    return rows


def archive(out, name):
    """Zip the staged tree, because Zenodo's uploader discards directories.

    Files dragged into the web uploader arrive as a flat list, so `OC2/` and
    `figures/` vanish and every name has to be unique on its own. A zip is the
    only way to hand over a structure, and it is what the manifest's paths
    describe.

    README.md and MANIFEST.csv go inside it and stay outside it as well. Inside
    so that someone who downloads only the zip has the explanation; outside so
    the Zenodo record itself shows them without anyone downloading anything,
    which a record whose sole file is an archive does not.
    """
    zip_path = out / name
    if zip_path.exists():
        zip_path.unlink()
    # AppleDouble sidecars and .DS_Store never travel. The staged tree should
    # not contain them, but Finder writes them into any directory it is asked
    # to look at, and a deposit is not the place to find out.
    members = sorted(path for path in out.rglob("*")
                     if path.is_file() and path != zip_path
                     and not path.name.startswith("._")
                     and path.name != ".DS_Store"
                     and "__MACOSX" not in path.parts)
    with zipfile.ZipFile(str(zip_path), "w", zipfile.ZIP_DEFLATED) as archive_file:
        for path in members:
            archive_file.write(str(path), str(path.relative_to(out)))
    return zip_path, len(members)


def wrap_markdown(text, width=78):
    """Re-wrap prose after substitution, leaving structure alone.

    Link text and version numbers change length, so the template cannot be
    wrapped by hand and stay wrapped.  Fenced code, table rows, headings, and
    link definitions pass through untouched; paragraphs and list items are
    re-flowed, with a list item's continuation lines hanging under its text.
    """
    out, block = [], []
    fenced = False
    first = indent = ""

    def flush():
        if block:
            joined = " ".join(part.strip() for part in block)
            out.extend(textwrap.wrap(joined, width=width,
                                     initial_indent=first,
                                     subsequent_indent=indent,
                                     break_long_words=False,
                                     break_on_hyphens=False))
            del block[:]

    for line in text.split("\n"):
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            fenced = not fenced
            out.append(line)
            continue
        if fenced or not stripped or stripped.startswith(("#", "|", ">")) \
                or re.match(r"^\[[^\]]+\]:", stripped):
            flush()
            out.append(line)
            continue
        bullet = re.match(r"^(\s*)([-*]|\d+\.)\s+", line)
        if bullet:
            flush()
            first = bullet.group(1) + bullet.group(2) + " "
            indent = " " * len(first)
            block.append(line[bullet.end():])
            continue
        if not block:
            leading = re.match(r"^(\s*)", line).group(1)
            first = indent = leading
        block.append(line)
    flush()
    return "\n".join(out)


def software_line(version):
    if SOFTWARE_DOI:
        return "mLOS {0}, [{1}][tool]".format(version, SOFTWARE_DOI)
    return ("mLOS {0} (<SOFTWARE CONCEPT DOI -- fill in SOFTWARE_DOI in "
            "tools/make_deposit.py once the release is archived>)".format(
                version))


def results_readme(runs, rows, out):
    version = runs[0]["version"]
    lines = [
        "# mLOS results: length-of-stay analysis of Orange County shelter data",
        "",
        "Every file a full mLOS run writes, for each of the data sets below, "
        "produced by {0}. Nothing here was edited by hand.".format(
            software_line(version)),
        "",
        "The input each run read is not repeated here. It is published "
        "separately, as prepared data, at [{0}][input], and the preparation "
        "log that pins it by digest travels with each run below.".format(
            PREPARED_DOI),
        "",
        "## What is in it",
        "",
        "One folder per data set.",
        "",
        "| folder | data set | how far it has been checked |",
        "|---|---|---|",
    ]
    for run in runs:
        what, checked = DATASET[run["label"]]
        lines.append("| `{0}/` | {1} | {2} |".format(
            run["label"], what, checked))
    lines += [
        "",
        "Inside a folder:",
        "",
        "| file | contents |",
        "|---|---|",
        "| `analysis_results.xlsx` | the consolidated workbook: every table "
        "the run produced |",
        "| `results.json` | every computed number, machine-readable; the "
        "workbook is a rendering of this |",
        "| `*.png`, and a `*.csv` beside most of them | each plot, and the "
        "values of the curves it draws. A stacked bar plot has no CSV: its "
        "bars bin a curve rather than redraw one |",
        "| `analysis_log.txt` | the console log of the run |",
        "| `data_preparation_stats.csv` | how the input rows were "
        "transformed, removed, or kept |",
        "| `*_settings.yaml` | the settings file that produced all of it |",
        "| `*_data.md` | where the input came from, and what it covers |",
        "| `*_data_run.txt` | the preparation log: which ShelterDataPrep "
        "version made the input, and its digest |",
        "",
        "`MANIFEST.csv` lists every file with its SHA-256, its size, the data "
        "set it belongs to, the mLOS version and timestamp of the run, and "
        "the input file and digest that run read.",
        "",
        "## Reading a result",
        "",
        "Start with `analysis_results.xlsx`. Its cover sheet carries the run "
        "metadata, and every other sheet is a table of numbers rather than "
        "formatted text, so a cell can be read into anything.",
        "",
        "The two guides that document every sheet, plot, and estimator ship "
        "with the software rather than here; see the archive named above.",
        "",
        "## Rebuilding it",
        "",
        "1. Fetch the prepared input from [{0}][input] and put "
        "`<SET>_data.csv` in `data/`.".format(PREPARED_DOI),
        "2. Install mLOS at the version named above. The archive holds the "
        "whole repository tree; R, plus the `survival`, `yaml`, `jsonlite`, "
        "`openxlsx`, `flexsurv`, and `digest` packages, is what it needs.",
        "3. Copy the settings file from this deposit into `data/` and run:",
        "",
        "   ```",
        "   Rscript mlos_run_complete.R --settings data/<SET>_settings.yaml "
        "--data data/<SET>_data.csv --results results",
        "   ```",
        "",
        "Plots are rendered, so a PNG will not match byte for byte across "
        "machines. The CSVs beside them, the workbook's numbers, and "
        "`results.json` will.",
        "",
        "## How to cite",
        "",
        "Cite this deposit by its version DOI, and cite alongside it:",
        "",
        "- the prepared input, [{0}][input];".format(PREPARED_DOI),
        "- {0}.".format(software_line(version)),
        "",
        "## License",
        "",
        "CC BY 4.0, matching the prepared data this is derived from.",
        "",
        "[input]: https://doi.org/{0}".format(PREPARED_DOI),
    ]
    if SOFTWARE_DOI:
        lines.append("[tool]: https://doi.org/{0}".format(SOFTWARE_DOI))
    (out / "README.md").write_text(
        wrap_markdown("\n".join(lines)) + "\n", encoding="utf-8")


def deck_readme(runs, rows, out):
    version = runs[0]["version"]
    source = ("the results deposit, [{0}][results]".format(RESULTS_DOI)
              if RESULTS_DOI else
              "the results deposit (<RESULTS VERSION DOI -- fill in "
              "RESULTS_DOI in tools/make_deposit.py>)")
    lines = [
        "# mLOS slide decks: length-of-stay analysis of Orange County "
        "shelter data",
        "",
        "Generated PowerPoint decks and the workbook of every table each one "
        "made, built by the deck builder shipped with {0} from {1}.".format(
            software_line(version), source),
        "",
        "> **The deck builder is experimental and may stay that way.** Which "
        "findings reach a slide is decided by fixed rules against fixed "
        "thresholds, not by judgment. A run in which nothing much happened "
        "still fills its slides, and results a reader would care about, "
        "including ones plainly visible in the workbook and the plots, may be "
        "absent from the deck. Absence from a slide is not evidence that "
        "there is nothing there. Read a deck as raw material for a briefing, "
        "never as the record. The record is the results deposit.",
        "",
        "Every number on every slide was computed and written out by the "
        "analysis. The deck builder redoes none of the statistics.",
        "",
        "## What is in it",
        "",
        "One folder per data set.",
        "",
        "| folder | data set |",
        "|---|---|",
    ]
    for run in runs:
        lines.append("| `{0}/` | {1} |".format(run["label"],
                                               DATASET[run["label"]][0]))
    lines += [
        "",
        "Inside a folder: the deck, the workbook holding every table it "
        "built, `figures/`, the images it placed, and `<deck>_slides.json`, "
        "which names every slide's position, title, layout and run and is "
        "what an outline is written against. Beside them, "
        "`educational.pptx` teaches the reading of the curves on the same "
        "run, with its own `educational_figures/` and no workbook: its tables "
        "are the deck's. `OC_deck_settings.yaml` at the top level is the "
        "settings file all of them were built under.",
        "",
        "`MANIFEST.csv` lists every file with its SHA-256, its size, the data "
        "set it belongs to, and the mLOS version and run timestamp of the "
        "analysis behind it.",
        "",
        "## Rebuilding it",
        "",
        "Restore a run from the results deposit into `results/`, then from "
        "the repository root:",
        "",
        "```",
        "python3 -m mlos_review.deck",
        "python3 -m mlos_review.variant data/educational.md",
        "```",
        "",
        "Python 3.9 or newer, with the packages `pyproject.toml` declares.",
        "",
        "## How to cite",
        "",
        "Cite the results rather than the deck wherever the numbers are what "
        "matter. Cite this deposit only for the slides themselves, and name "
        "{0} alongside it.".format(source),
        "",
        "## License",
        "",
        "CC BY 4.0.",
    ]
    if RESULTS_DOI:
        lines.append("")
        lines.append("[results]: https://doi.org/{0}".format(RESULTS_DOI))
    if SOFTWARE_DOI:
        lines.append("[tool]: https://doi.org/{0}".format(SOFTWARE_DOI))
    (out / "README.md").write_text(
        wrap_markdown("\n".join(lines)) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--deck", action="store_true",
                        help="stage the deck bundle instead of the results")
    parser.add_argument("destination", nargs="?", default=None)
    args = parser.parse_args()

    runs = read_runs()
    for run in runs:
        check_input(run)

    # Named for the sets it holds and the day it was made, matching the file
    # published at 0.1.0. The date rather than the version is a deliberate
    # choice of the maintainer's; the version is inside, in every run log and
    # in the README's citation block.
    sets = "_".join(run["label"] for run in runs)
    stamp = datetime.date.today().isoformat()
    if args.deck:
        out = Path(args.destination) if args.destination else DECK_OUT
        rows = stage_deck(runs, out)
        deck_readme(runs, rows, out)
        name = "mlos_deck_{0}_{1}.zip".format(sets, stamp)
        if RESULTS_DOI is None:
            print("note: RESULTS_DOI is unset, so the README carries a "
                  "placeholder; set it and re-run once the results deposit "
                  "is published")
    else:
        out = Path(args.destination) if args.destination else RESULTS_OUT
        rows = stage_results(runs, out)
        results_readme(runs, rows, out)
        name = "mlos_results_{0}_{1}.zip".format(sets, stamp)

    if SOFTWARE_DOI is None:
        print("note: SOFTWARE_DOI is unset, so the README carries a "
              "placeholder; set it and re-run once the release is archived")

    zip_path, zipped = archive(out, name)
    total = sum(int(row[6]) for row in rows)
    print("{0}: {1} files, {2:.1f} MB, mLOS {3}".format(
        out, len(rows) + 2, total / 1e6, runs[0]["version"]))
    for run in runs:
        print("  {0}: {1}, prepared by shelterprep {2}".format(
            run["label"], run["data_file"], run["prep_version"]))
    print("\nUpload these three, and nothing else:")
    for path in (zip_path, out / "README.md", out / "MANIFEST.csv"):
        print("  {0}  ({1:.1f} MB)".format(path.name, path.stat().st_size / 1e6))
    print("The zip holds all {0} files with their directories, which Zenodo's "
          "uploader discards otherwise. The other two ride outside it so the "
          "record shows them without a download.".format(zipped))


if __name__ == "__main__":
    main()
